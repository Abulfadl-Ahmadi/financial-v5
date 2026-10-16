"""Read-only reconciliation of supplied bank statements and the site's local SQLite snapshot.

Run with the bundled Python/openpyxl. Never writes to backend/db.sqlite3.
Classification is evidence of a known payee, NOT proof of payroll or a business expense.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
import re
import sqlite3
import unicodedata
import warnings

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
DB = HERE.parent / 'backend' / 'db.sqlite3'
FILES = {
    '059a522e-67c3-4a95-95b2-4b99c02dc220.xlsx': '873-800-4851219-1',
    '06da3bf2-161d-4df1-803d-93d31df29936.xlsx': '2137-888-4851219-1',
}
SOURCE_HEADER = ('ردیف', 'تاریخ', 'شرح سند', 'شماره سند', 'شماره برگه/چک',
                 'شناسه واریز', 'شعبه', 'واریز (ریال)', 'برداشت (ریال)',
                 'مانده (ریال)', 'توضیحات کاربر')
HEADERS = ['حساب بانکی', 'فایل اصلی', 'ردیف بانک', 'سطر اکسل اصلی', 'تاریخ و ساعت شمسی',
           'نوع', 'واریز (ریال)', 'برداشت (ریال)', 'مانده (ریال)', 'شماره سند',
           'شناسه واریز', 'کد پیگیری کارت', 'کارت مقصد', 'شناسه حساب سایت',
           'نام صاحب کارت در سایت', 'نقش حساب/کاربر', 'شناسه فیش سایت',
           'وضعیت تطبیق فیش', 'گروه', 'دلیل و اقدام', 'شرح کامل بانک', 'توضیح کاربر']


def digits(value):
    return ''.join(str(unicodedata.digit(c)) if c.isdigit() else c for c in str(value or ''))


def amount(value):
    return int(digits(value).replace(',', '').replace('٬', '').replace(' ', '') or '0')


def safe(value):
    """Keep text as data in Excel, rather than executing untrusted bank text as formulas."""
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


# URI mode=ro prevents even accidental database writes.
connection = sqlite3.connect(DB.resolve().as_uri() + '?mode=ro', uri=True)
connection.row_factory = sqlite3.Row
accounts = {}
for item in connection.execute('''
    SELECT a.id, a.f_name, a.l_name, a.card_number, a.is_workshop_account,
           a.user_id, u.role, u.first_name AS user_first, u.last_name AS user_last
    FROM financial_account AS a LEFT JOIN users_user AS u ON u.id=a.user_id
'''):
    accounts[item['card_number']] = dict(item)
receipts = [dict(item) for item in connection.execute('''
    SELECT id, tracking_code, date_jalali, amount, from_account_id, to_account_id, status, receipt_type
    FROM financial_receipt
''')]
connection.close()

warnings.filterwarnings('ignore', message='Workbook contains no default style')
rows, controls = [], []
for filename, account in FILES.items():
    ws = openpyxl.load_workbook(HERE / filename, read_only=True, data_only=True).active
    assert ws.title == 'Report'
    iterator = ws.iter_rows(values_only=True)
    meta = [next(iterator) for _ in range(4)]
    assert tuple(meta[3]) == SOURCE_HEADER, filename
    assert account.replace('-', '') in re.sub(r'\D', '', digits(meta[0][0]))
    local = []
    for excel_row, raw in enumerate(iterator, 5):
        if raw[0] is None:
            continue
        desc = digits(raw[2]); note = str(raw[10] or '').strip()
        destination = re.search(r'به ک\s*([0-9]{16})(?!\d)', desc)
        sender = re.search(r'از ک\s*([0-9]{16})(?!\d)', desc)
        tracking = re.search(r'ش م\s*(\d+)', desc)
        item = dict(account=account, filename=filename, number=int(raw[0]), excel_row=excel_row,
                    date=digits(raw[1]), deposit=amount(raw[7]), withdrawal=amount(raw[8]),
                    balance=amount(raw[9]), document=digits(raw[3]), reference=digits(raw[5]),
                    tracking=tracking.group(1) if tracking else '',
                    destination=destination.group(1) if destination else '',
                    sender=sender.group(1) if sender else '', desc=desc, note=note,
                    site_id='', site_name='', role='', receipt_id='', receipt_state='',
                    category='', reason='')
        assert not (item['deposit'] and item['withdrawal']), (filename, excel_row)
        local.append(item)
    assert [x['number'] for x in local] == list(range(1, len(local)+1)), filename
    declared_in = amount(re.search(r'[\d,]+', digits(meta[1][8])).group())
    declared_out = amount(re.search(r'[\d,]+', digits(meta[2][8])).group())
    incoming = sum(x['deposit'] for x in local)
    outgoing = sum(x['withdrawal'] for x in local)
    assert (incoming, outgoing) == (declared_in, declared_out), filename
    controls.append([account, filename, len(local), incoming, declared_in, outgoing, declared_out, 'تطبیق دارد'])
    rows.extend(local)

# A bank transfer between the two provided deposits must be counted once, as neither income nor expense.
internal = defaultdict(list)
for row in rows:
    if 'انتقال وجه موبایلت' in row['desc'] and '4851219' in row['desc'] and row['reference']:
        internal[row['reference']].append(row)
internal_pairs = []
for ref, pair in internal.items():
    if (len(pair) == 2 and len({x['account'] for x in pair}) == 2
            and sum(x['deposit'] for x in pair) == sum(x['withdrawal'] for x in pair)
            and sum(x['deposit'] for x in pair) > 0):
        internal_pairs.append(pair)
        for row in pair:
            row['category'] = 'انتقال داخلی'
            row['reason'] = 'هر دو طرف انتقال بین سپرده‌های ارائه‌شده با شناسه و مبلغ متناظر؛ درآمد/هزینه نیست.'

# Infer source-card-to-deposit mapping from fee rows, which contain the source card.
# The transfer row itself normally has only a destination card ("به ک"), not a source.
source_deposit = defaultdict(set)
for row in rows:
    if row['sender'] and 'کارمزد' in row['desc']:
        source_deposit[row['sender']].add(row['account'])
# Exact reconciliation requires tracking, day, amount, source deposit and destination card.
# Receipt.amount in site is TOMAN; statements are RIAL.
by_tracking = defaultdict(list)
for row in rows:
    if row['tracking'] and row['withdrawal']:
        by_tracking[row['tracking']].append(row)
receipt_results = []
for receipt in receipts:
    source = next((a['card_number'] for a in accounts.values() if a['id'] == receipt['from_account_id']), '')
    dest = next((a['card_number'] for a in accounts.values() if a['id'] == receipt['to_account_id']), '')
    valid = lambda row: (row['date'][:10] == digits(receipt['date_jalali'])
                         and row['withdrawal'] == receipt['amount'] * 10
                         and row['destination'] == dest
                         and row['account'] in source_deposit.get(source, set()))
    exact = [x for x in by_tracking[digits(receipt['tracking_code'])] if valid(x)]
    probable = [] if exact else [x for x in rows if valid(x)]
    if len(exact) == 1:
        match, state = exact[0], 'تطبیق دقیق'
    elif not exact and len(probable) == 1:
        match, state = probable[0], 'تطبیق محتمل؛ کد پیگیری متفاوت'
    else:
        match, state = None, 'بدون تطبیق یکتا'
    if match:
        match['receipt_id'] = receipt['id']
        match['receipt_state'] = state
    receipt_results.append([receipt['id'], receipt['date_jalali'], receipt['tracking_code'],
                            receipt['amount'], receipt['amount'] * 10, receipt['from_account_id'],
                            receipt['to_account_id'], dest, receipt['status'], receipt['receipt_type'],
                            state, match['account'] if match else '', match['number'] if match else '',
                            match['tracking'] if match else '',
                            'در صورت تفاوت کد، شماره پیگیری را با رسید اصلی اصلاح کنید؛ دوباره ثبت نکنید' if probable else ''])

for row in rows:
    if row['category']:
        continue
    dest = accounts.get(row['destination'])
    if dest:
        row['site_id'] = dest['id']
        row['site_name'] = (dest['f_name'] + ' ' + dest['l_name']).strip()
        row['role'] = dest['role'] or ('حساب کارگاه' if dest['is_workshop_account'] else 'بدون کاربر')
    note, desc = row['note'], row['desc']
    if 'برگشت خورد' in note or 'برگشت خورده' in note:
        row['category'] = 'معلق: ادعای برگشت'
        row['reason'] = 'تا مشاهده واریز برگشتی/رسید و تعیین وضعیت اولیه، هزینه قطعی منظور نشود.'
    elif 'فکنم' in note or 'مطمئن نیستم' in note:
        row['category'] = 'معلق: تردید صریح درباره دوخت'
        row['reason'] = 'توضیح کاربر نسبت به ارتباط با دوخت نامطمئن است؛ نیازمند تایید.'
    elif 'تعمیر چرخ' in note:
        row['category'] = 'تعمیر چرخ؛ تأیید لازم'
        row['reason'] = 'توضیح کاربر مرتبط با تعمیر چرخ است؛ باید معلوم شود چرخ خیاطی کارگاه بوده است.'
    elif dest and not dest['is_workshop_account'] and dest['role'] in ('WORKER', 'IRONER', 'SEWING_SUPERVISOR', 'CUTTING_SUPERVISOR') and row['withdrawal']:
        row['category'] = 'مقصد پرسنل؛ هدف پرداخت تأیید نشده'
        row['reason'] = 'شماره کارت دقیقاً در حساب‌های سایت و کاربر با نقش کاری ثبت شده؛ اجرت/حقوق بودن از بانک اثبات نمی‌شود.'
    elif 'بابایی' in desc:
        row['category'] = 'نام مشابه؛ بررسی هویت'
        row['reason'] = 'اشتراک نام خانوادگی اثبات ارتباط با سفارش/کارفرما نیست.'
    elif any(x in note for x in ('چال ماهی', 'تمدید اقامت', 'برای قبض')) or note == 'تاکسی':
        row['category'] = 'توضیح غیرمرتبط یا نامعلوم'
        row['reason'] = 'مدرک ارتباط با کارگاه موجود نیست؛ برای پرداخت کاری احتمالی رسید ارائه شود.'
    elif 'کارمزد' in desc:
        row['category'] = 'کارمزد؛ تعیین تکلیف تراکنش پایه'
        row['reason'] = 'با شماره سند/شناسه به تراکنش پایه وصل شود؛ به‌تنهایی هزینه کارگاه نیست.'
    elif 'سودعلی الحساب' in desc:
        row['category'] = 'سود حساب؛ ماهیت نامعلوم'
        row['reason'] = 'سهم درآمد کاری/شخصی نامعلوم است.'
    elif ' اصلاح سند' in desc:
        row['category'] = 'اصلاح سند؛ بررسی جفت'
        row['reason'] = 'برای پیشگیری از شمارش مضاعف باید به سند اصلاحی/اولیه متصل شود.'
    elif 'خرید کالا و خدمات' in desc:
        row['category'] = 'خرید POS؛ بررسی فاکتور'
        row['reason'] = 'شرح بانک نوع کالا/هویت فروشنده را مشخص نمی‌کند.'
    elif 'برداشت وجه از خودپرداز' in desc:
        row['category'] = 'نقدی؛ بررسی مصرف'
        row['reason'] = 'مصرف پول نقد بدون سند تعیین نشده است.'
    elif row['withdrawal'] and (row['destination'] or 'انتقال وجه آنی به شبا' in desc or 'انتقال وجه موبایلت به سپرده' in desc):
        row['category'] = 'خروجی؛ بررسی ذی‌نفع'
        row['reason'] = 'ارتباط ذی‌نفع با کارگاه در داده‌های فعلی اثبات نشد.'
    elif row['deposit']:
        row['category'] = 'واریز؛ بررسی منشأ'
        row['reason'] = 'درآمد سفارش، قرض یا انتقال شخصی بدون سند تفکیک نمی‌شود.'
    else:
        row['category'] = 'سایر؛ بررسی'
        row['reason'] = 'داده برای انتساب به هزینه کارگاه کافی نیست.'

# Link fees to base transaction by account + exact bank document + exact bank reference.
base_by_key = defaultdict(list)
for row in rows:
    if not row['category'].startswith('کارمزد') and row['document'] and row['reference']:
        base_by_key[(row['account'],row['document'],row['reference'])].append(row)
for row in rows:
    if not row['category'].startswith('کارمزد') or not row['document'] or not row['reference']:
        continue
    bases = [x for x in base_by_key[(row['account'],row['document'],row['reference'])]
             if x['date'][:10] == row['date'][:10] and x['withdrawal'] and x is not row]
    if len(bases) == 1:
        base = bases[0]
        row['reason'] = f"تراکنش پایه ردیف بانک {base['number']} ({base['category']})؛ فقط با تایید ماهیت آن تخصیص دهید."


def values(row):
    return [row['account'], row['filename'], row['number'], row['excel_row'], row['date'],
            'واریز' if row['deposit'] else 'برداشت', row['deposit'], row['withdrawal'],
            row['balance'], row['document'], row['reference'], row['tracking'],
            row['destination'], row['site_id'], row['site_name'], row['role'],
            row['receipt_id'], row['receipt_state'], row['category'], row['reason'],
            row['desc'], row['note']]


workbook = openpyxl.Workbook()
workbook.remove(workbook.active)

def sheet(title, headers, entries):
    ws = workbook.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for entry in entries:
        ws.append([safe(value) for value in entry])
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='18375C')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
    for col in range(1, len(headers)+1):
        ws.column_dimensions[get_column_letter(col)].width = 58 if headers[col-1] in ('شرح کامل بانک','دلیل و اقدام') else 29
    for line in ws.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, int):
                cell.number_format = '#,##0'
    return ws

sheet('کنترل جمع بانک', ['حساب','فایل','تعداد','واریز محاسبه‌شده','واریز بانک','برداشت محاسبه‌شده','برداشت بانک','نتیجه'], controls)
known = [x for x in rows if x['category'] in ('مقصد پرسنل؛ هدف پرداخت تأیید نشده','تعمیر چرخ؛ تأیید لازم')]
sheet('انتقال به پرسنل', HEADERS, (values(x) for x in known))
sheet('فیش‌های تطبیق‌شده', HEADERS, (values(x) for x in rows if x['receipt_id']))
sheet('کنترل فیش‌های سایت', ['شناسه فیش','تاریخ','پیگیری سایت','مبلغ تومان','مبلغ ریال','شناسه مبدأ','شناسه مقصد','کارت مقصد','وضعیت سایت','نوع فیش سایت','وضعیت تطبیق','حساب بانک','ردیف بانک','پیگیری بانک','اقدام'], receipt_results)
sheet('موارد حساس', HEADERS, (values(x) for x in rows if x['category'].startswith('معلق:') or x['category'] in ('تعمیر چرخ؛ تأیید لازم','نام مشابه؛ بررسی هویت')))
sheet('باقی برای بازبینی', HEADERS, (values(x) for x in rows if x['category'] not in ('انتقال داخلی','مقصد پرسنل؛ هدف پرداخت تأیید نشده','تعمیر چرخ؛ تأیید لازم')))
sheet('انتقال داخلی', HEADERS, (values(x) for x in rows if x['category'] == 'انتقال داخلی'))
sheet('تمام تراکنش‌ها', HEADERS, (values(x) for x in rows))
per_card = defaultdict(list)
for x in known:
    per_card[x['destination']].append(x)
sheet('خلاصه کارت پرسنل', ['شناسه حساب سایت','نام ثبت‌شده','نقش کاربر','کارت','تعداد','جمع ریال','تطبیق دقیق فیش','تطبیق محتمل فیش','توضیح'],
      ([xs[0]['site_id'],xs[0]['site_name'],xs[0]['role'],card,len(xs),sum(x['withdrawal'] for x in xs),
        sum(x['receipt_state']=='تطبیق دقیق' for x in xs),sum(x['receipt_state'].startswith('تطبیق محتمل') for x in xs),
        'پرداخت به کارت ثبت‌شده؛ ماهیت اجرت/قرض/شخصی را تأیید کنید']
       for card,xs in sorted(per_card.items(),key=lambda pair:-sum(x['withdrawal'] for x in pair[1]))))
sheet('راهنما', ['موضوع','شرح'], [
    ['واحد پول','فایل بانک ریال است؛ فیش و هزینه سایت تومان‌اند (تومان = ریال ÷ ۱۰).'],
    ['سطر بانک','ردیف اکسل اصلی = ردیف بانک + ۴؛ فایل و حساب در هر سطر درج شده است.'],
    ['معنی ارتباط','شماره کارت ثبت‌شده برای کاربر کارگاه نشانه ذی‌نفع است، نه اثبات حقوق یا هزینه کارگاه.'],
    ['فیش قبلی','تطبیق دقیق: کد پیگیری، حساب مبدأ/مقصد، روز و مبلغ؛ تطبیق محتمل: همه به‌جز کد پیگیری. هیچ‌کدام دوباره وارد نشوند.'],
    ['واریز برگشتی','برگشت ادعایی تا اثبات جفت بانکی آن معلق است.'],
    ['انتقال داخلی','دو طرف انتقال بین دو سپرده در گزارش درآمد/هزینه محسوب نشوند.'],
    ['حریم خصوصی','صورت‌حساب‌ها شامل اطلاعات حساس بانکی هستند؛ این خروجی را خصوصی نگه دارید.'],
])
output = HERE / 'تراکنش‌های-کارگاه-برای-بررسی.xlsx'
workbook.save(output)
print('OUTPUT', output)
print('BANK CONTROL', controls)
print('CATEGORIES', Counter(x['category'] for x in rows))
print('KNOWN-PAYEE',len(known),sum(x['withdrawal'] for x in known))
print('SITE RECEIPTS',Counter(x[10] for x in receipt_results))
print('INTERNAL PAIRS',len(internal_pairs))
